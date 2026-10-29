import contextlib
import io
import os
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET

import csv2xlsx

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def run(*argv):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = csv2xlsx.main(list(argv))
    return code, out.getvalue(), err.getvalue()


def read_cells(path):
    """Read {ref: (type, text)} from sheet1 without any third-party lib."""
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
    cells = {}
    for c in root.iter("{%s}c" % NS["m"]):
        if c.get("t") == "inlineStr":
            cells[c.get("r")] = ("s", "".join(t.text or "" for t in c.iter("{%s}t" % NS["m"])))
        else:
            cells[c.get("r")] = ("n", c.find("m:v", NS).text)
    return cells


class DecodeTests(unittest.TestCase):
    def test_utf8_bom(self):
        self.assertEqual(csv2xlsx.decode(b"\xef\xbb\xbfa,b"), ("a,b", "utf-8-sig"))

    def test_utf16(self):
        text, enc = csv2xlsx.decode("a,b".encode("utf-16"))
        self.assertEqual((text, enc), ("a,b", "utf-16"))

    def test_plain_utf8(self):
        self.assertEqual(csv2xlsx.decode("café".encode("utf-8"))[1], "utf-8")

    def test_cp1252_fallback(self):
        text, enc = csv2xlsx.decode("café".encode("cp1252"))
        self.assertEqual((text, enc), ("café", "cp1252"))

    def test_undefined_cp1252_byte_falls_to_latin1(self):
        self.assertEqual(csv2xlsx.decode(b"a\x81b")[1], "latin-1")

    def test_forced_bad_encoding(self):
        with self.assertRaises(csv2xlsx.CsvError):
            csv2xlsx.decode(b"\xff", "ascii")


class DelimiterTests(unittest.TestCase):
    def test_semicolon(self):
        self.assertEqual(csv2xlsx.detect_delimiter("a;b;c\n1;2;3\n"), ";")

    def test_tab(self):
        self.assertEqual(csv2xlsx.detect_delimiter("a\tb\n1\t2\n"), "\t")

    def test_comma_in_quotes_not_confused(self):
        self.assertEqual(csv2xlsx.detect_delimiter('n;note\n1;"a, b, c"\n2;"x, y"\n'), ";")

    def test_single_column_defaults_comma(self):
        self.assertEqual(csv2xlsx.detect_delimiter("a\nb\n"), ",")


class CellTests(unittest.TestCase):
    def test_numbers(self):
        self.assertEqual(csv2xlsx.cell_number("12", "."), "12")
        self.assertEqual(csv2xlsx.cell_number("-3.5", "."), "-3.5")

    def test_leading_zero_and_long_ids_stay_text(self):
        self.assertIsNone(csv2xlsx.cell_number("007", "."))
        self.assertIsNone(csv2xlsx.cell_number("1234567890123456", "."))

    def test_decimal_comma(self):
        self.assertEqual(csv2xlsx.cell_number("1,5", ","), "1.5")
        self.assertIsNone(csv2xlsx.cell_number("1,5", "."))

    def test_col_letter(self):
        self.assertEqual([csv2xlsx.col_letter(i) for i in (0, 25, 26, 27, 701, 702)],
                         ["A", "Z", "AA", "AB", "ZZ", "AAA"])


class CliTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def make(self, name, data):
        p = self.path(name)
        with open(p, "wb") as f:
            f.write(data)
        return p

    def test_converts_semicolon_cp1252_with_decimal_comma(self):
        src = self.make("sales.csv", "naam;bedrag\nJosé;1,5\nZoë;007\n".encode("cp1252"))
        code, out, _ = run(src, "--decimal", ",")
        self.assertEqual(code, 0)
        self.assertIn("cp1252", out)
        cells = read_cells(self.path("sales.xlsx"))
        self.assertEqual(cells["A2"], ("s", "José"))
        self.assertEqual(cells["B2"], ("n", "1.5"))
        self.assertEqual(cells["B3"], ("s", "007"))

    def test_roundtrip_with_openpyxl_if_present(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        src = self.make("t.csv", b'id,name,qty\n1,"a, b",2.5\n2,"line\nbreak",\n')
        self.assertEqual(run(src)[0], 0)
        ws = openpyxl.load_workbook(self.path("t.xlsx")).active
        self.assertEqual(ws["B2"].value, "a, b")
        self.assertEqual(ws["B3"].value, "line\nbreak")
        self.assertEqual(ws["C2"].value, 2.5)
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws.freeze_panes, "A2")

    def test_xml_special_chars_and_control_chars(self):
        src = self.make("x.csv", b'a,b\n"<&> ""q""",x\x01y\n')
        self.assertEqual(run(src)[0], 0)
        cells = read_cells(self.path("x.xlsx"))
        self.assertEqual(cells["A2"], ("s", '<&> "q"'))
        self.assertEqual(cells["B2"], ("s", "xy"))

    def test_missing_input(self):
        code, _, err = run(self.path("nope.csv"))
        self.assertEqual(code, 2)
        self.assertIn("cannot read", err)

    def test_empty_input(self):
        code, _, err = run(self.make("e.csv", b"  \n"))
        self.assertEqual(code, 2)
        self.assertIn("empty", err)

    def test_refuses_overwrite_then_force(self):
        src = self.make("a.csv", b"a,b\n1,2\n")
        self.assertEqual(run(src)[0], 0)
        code, _, err = run(src)
        self.assertEqual(code, 1)
        self.assertIn("exists", err)
        self.assertEqual(run(src, "--force")[0], 0)

    def test_bad_delimiter_arg(self):
        src = self.make("a.csv", b"a,b\n1,2\n")
        code, _, err = run(src, "--delimiter", "ab")
        self.assertEqual(code, 2)
        self.assertIn("single character", err)

    def test_forced_delimiter_tab_escape(self):
        src = self.make("a.csv", b"a\tb\n1\t2\n")
        self.assertEqual(run(src, "--delimiter", "\\t")[0], 0)
        self.assertEqual(read_cells(self.path("a.xlsx"))["B2"], ("n", "2"))

    def test_ragged_rows(self):
        src = self.make("r.csv", b"a,b,c\n1\n1,2,3,4\n")
        self.assertEqual(run(src)[0], 0)
        cells = read_cells(self.path("r.xlsx"))
        self.assertEqual(cells["D3"], ("n", "4"))

    def test_header_only(self):
        src = self.make("h.csv", b"a,b\n")
        self.assertEqual(run(src)[0], 0)

    def test_oversized_cell_truncated_with_warning(self):
        src = self.make("big.csv", b"a\n" + b"x" * 40000 + b"\n")
        code, _, err = run(src)
        self.assertEqual(code, 0)
        self.assertIn("warning", err)
        self.assertEqual(len(read_cells(self.path("big.xlsx"))["A2"][1]), 32767)

    def test_unwritable_output(self):
        src = self.make("a.csv", b"a,b\n1,2\n")
        code, _, err = run(src, "-o", os.path.join(self.tmp.name, "no_dir", "o.xlsx"))
        self.assertEqual(code, 1)
        self.assertIn("cannot write", err)


if __name__ == "__main__":
    unittest.main()
